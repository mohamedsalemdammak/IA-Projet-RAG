import json
import re
import unicodedata
from pathlib import Path
from PIL import Image

try:
    import pymupdf
except ImportError:
    pymupdf = None
try:
    import pytesseract
except ImportError:
    pytesseract = None
try:
    import docx
except ImportError:
    docx = None
try:
    from pptx import Presentation
except ImportError:
    Presentation = None
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from config import OCR_LANG

DATE_PATTERN = re.compile(
    r'\b(?:\d{1,2}\s+(?:janvier|février|fevrier|mars|avril|mai|juin|juillet|'
    r'août|aout|septembre|octobre|novembre|décembre|decembre|جانفي|فيفري|'
    r'مارس|أفريل|افريل|ماي|جوان|جويلية|أوت|اوت|سبتمبر|أكتوبر|اكتوبر|'
    r'نوفمبر|ديسمبر|يناير|فبراير|أبريل|ابريل|يونيو|يوليو|أغسطس|اغسطس)\s+\d{4}'
    r'|\d{1,2}[./-]\d{1,2}[./-]\d{2,4}'
    r'|\d{4}[./-]\d{1,2}[./-]\d{1,2})\b',
    re.IGNORECASE,
)
TOKEN_PATTERN = re.compile(r'\w+|[^\w\s]', re.UNICODE)


def check_ocr():
    if pytesseract is None:
        return False
    try:
        version = pytesseract.get_tesseract_version()
        languages = set(pytesseract.get_languages(config=''))
        required = {'ara', 'eng', 'fra'}
        print(f'OK Tesseract: {version}')
        missing = required - languages
        if missing:
            print(f'WARNING OCR languages missing: {missing}')
            return False
        return True
    except Exception as exc:
        print(f'WARNING Tesseract unavailable: {exc}')
        return False


def clean_text(text):
    if not text:
        return ''
    text = text.replace('\x00', ' ')
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n\s*\n\s*\n+', '\n\n', text)
    return text.strip()


def count_tokens(text):
    return len(TOKEN_PATTERN.findall(text or ''))


def _normalize_label(text):
    normalized = unicodedata.normalize('NFKD', text).casefold()
    normalized = ''.join(char for char in normalized if not unicodedata.combining(char))
    return re.sub(r'[\s_-]+', ' ', normalized).strip()


def _get_document_type(relative_path):
    folders = [
        _normalize_label(part)
        for part in relative_path.replace('\\', '/').split('/')[:-1]
    ]
    type_by_folder = {
        'pv': 'pv',
        'rapport moral': 'rm',
        'rapport financiere': 'rf',
    }
    return next((type_by_folder[folder] for folder in folders if folder in type_by_folder), None)


def _extract_title(text, filename):
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for line in lines[:30]:
        normalized = _normalize_label(line)
        if re.search(r'\b(?:pv|proces verbal|rapport(?: moral| financier)?)\b', normalized):
            return line[:200]
        if re.match(r'^(?:titre|title)\s*[:\-]', normalized):
            return re.split(r'[:\-]', line, maxsplit=1)[-1].strip()[:200] or Path(filename).stem
    return Path(filename).stem


def _extract_place(text):
    lines = [line.strip() for line in text.splitlines()]
    place_label = re.compile(
        r'^\s*(?:lieu|place|location|المكان|مكان الاجتماع|مكان الانعقاد|'
        r'مكان التنفيذ|مكان|مقر)\s*(?:[:：-]\s*(.*))?$',
        re.IGNORECASE,
    )
    field_label = re.compile(r'^\s*(?:date|heure|time|date et heure)\b', re.IGNORECASE)
    time_value = re.compile(r'^\d{1,2}(?:h|:)\d{0,2}\s*(?:à\s*\d{1,2}(?:h|:)\d{0,2})?$', re.IGNORECASE)

    for index, line in enumerate(lines):
        match = place_label.match(line)
        if not match:
            continue
        if match.group(1):
            return match.group(1).strip()[:200] or None
        for candidate in lines[index + 1:index + 5]:
            if not candidate or field_label.match(candidate):
                continue
            if DATE_PATTERN.search(candidate) or time_value.match(candidate):
                continue
            return candidate[:200]
    return None


def extract_document_metadata(file_path, relative_path, text):
    date_match = DATE_PATTERN.search(text)
    return {
        'title': _extract_title(text, file_path.name),
        'type': _get_document_type(relative_path),
        'date': date_match.group(0) if date_match else None,
        'place': _extract_place(text),
        'document_name': file_path.name,
        'document_token_count': count_tokens(text),
    }


def ocr_image(image):
    if pytesseract is None:
        return ''
    try:
        return clean_text(pytesseract.image_to_string(image, lang=OCR_LANG))
    except Exception as exc:
        print(f'WARNING OCR error: {exc}')
        return ''


def extract_pdf(file_path, ocr_available=True):
    if pymupdf is None:
        raise RuntimeError('PyMuPDF is not installed')
    results = []
    doc = pymupdf.open(file_path)
    for page_number, page in enumerate(doc, start=1):
        text = clean_text(page.get_text('text'))
        used_ocr = False
        if len(text) < 50 and ocr_available:
            pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            image = Image.frombytes('RGB', [pix.width, pix.height], pix.samples)
            ocr_text = ocr_image(image)
            if len(ocr_text) > len(text):
                text, used_ocr = ocr_text, True
        if text:
            results.append({'text': text, 'page': page_number, 'ocr': used_ocr})
    doc.close()
    return results


def extract_docx(file_path):
    if docx is None:
        raise RuntimeError('python-docx is not installed')
    texts = []
    doc = docx.Document(file_path)
    for paragraph in doc.paragraphs:
        text = clean_text(paragraph.text)
        if text:
            texts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells]
            row_text = ' | '.join(x for x in cells if x)
            if row_text:
                texts.append(row_text)
    return [{'text': t, 'page': None, 'ocr': False} for t in texts]


def extract_pptx(file_path):
    if Presentation is None:
        raise RuntimeError('python-pptx is not installed')
    results = []
    presentation = Presentation(file_path)
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, 'text'):
                text = clean_text(shape.text)
                if text:
                    texts.append(text)
        if texts:
            results.append({'text': '\n'.join(texts), 'slide': slide_number, 'ocr': False})
    return results


def extract_xlsx(file_path):
    if load_workbook is None:
        raise RuntimeError('openpyxl is not installed')
    results = []
    wb = load_workbook(file_path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            values = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if values:
                rows.append(' | '.join(values))
        if rows:
            results.append({'text': '\n'.join(rows), 'sheet': ws.title, 'ocr': False})
    wb.close()
    return results


def extract_image(file_path):
    with Image.open(file_path) as image:
        text = ocr_image(image)
    return [{'text': text, 'ocr': bool(text)}] if text else []


def extract_plain(file_path):
    text = clean_text(Path(file_path).read_text(encoding='utf-8', errors='ignore'))
    return [{'text': text, 'ocr': False}] if text else []


def extract_file(file_path, ocr_available=True):
    ext = Path(file_path).suffix.lower()
    if ext == '.pdf': return extract_pdf(file_path, ocr_available)
    if ext == '.docx': return extract_docx(file_path)
    if ext == '.pptx': return extract_pptx(file_path)
    if ext in {'.xlsx', '.xlsm'}: return extract_xlsx(file_path)
    if ext in {'.png', '.jpg', '.jpeg', '.webp', '.tiff', '.bmp'}: return extract_image(file_path)
    if ext in {'.txt', '.md', '.py', '.json', '.csv'}: return extract_plain(file_path)
    return None


def ingest_folder(input_dir, output_report):
    input_dir = Path(input_dir)
    file_report = []
    failed, unsupported, empty, ocr_files = [], [], [], []
    ocr_available = check_ocr()
    files = [p for p in input_dir.rglob('*') if p.is_file()]

    for file_path in files:
        relative = str(file_path.relative_to(input_dir))
        document_type = _get_document_type(relative)
        try:
            sections = extract_file(file_path, ocr_available)
            if sections is None:
                unsupported.append(relative)
                file_report.append({
                    'file': relative, 'type': document_type,
                    'status': 'UNSUPPORTED',
                })
                continue
            document_text = clean_text('\n'.join(section.get('text', '') for section in sections))
            document_metadata = extract_document_metadata(file_path, relative, document_text)
            file_used_ocr = False
            for section in sections:
                file_used_ocr |= bool(section.get('ocr', False))
            if file_used_ocr:
                ocr_files.append(relative)
            if not document_text:
                empty.append(relative)
                status = 'EMPTY'
            else:
                status = 'OCR' if file_used_ocr else 'SUCCESS'
            file_report.append({
                'file': relative,
                'status': status,
                'ocr': file_used_ocr,
                'extracted_text': document_text,
                'extension': file_path.suffix.lower(),
                **document_metadata,
            })
        except Exception as exc:
            failed.append({'file': relative, 'error': str(exc)})
            file_report.append({
                'file': relative, 'type': document_type,
                'status': 'FAILED', 'error': str(exc),
            })

    output_report = Path(output_report); output_report.parent.mkdir(parents=True, exist_ok=True)
    report = {
        'folder': str(input_dir), 'total_files': len(files),
        'successful_files': sum(x['status'] in {'SUCCESS', 'OCR'} for x in file_report),
        'failed_files': len(failed), 'unsupported_files': len(unsupported),
        'empty_files': len(empty), 'ocr_files': len(ocr_files),
        'failed': failed,
        'unsupported': unsupported, 'empty': empty, 'ocr_files_list': ocr_files,
        'files': file_report,
    }
    output_report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report
